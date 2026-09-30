# -*- coding: utf-8 -*-
"""把汇报稿 Markdown 转换为排版好的 Word 文档。

约定：
  1. 首行 = 主标题，紧随的“（…）”行 = 副标题
  2. 独占一行的“一、二、三…” = 一级标题（章）
  3. 独占一行的“（一）（二）…” = 二级标题（节）
  4. 缩进的“一、二、三…” = 悬挂缩进的小点
"""
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

SRC = sys.argv[1]
DST = sys.argv[2]

doc = Document()

normal = doc.styles["Normal"]
normal.font.name = "宋体"
normal.font.size = Pt(12)
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
normal.paragraph_format.line_spacing = 1.5
normal.paragraph_format.space_after = Pt(0)

for section in doc.sections:
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.17)
    section.right_margin = Cm(3.17)

CN_NUM = "一二三四五六七八九十"
RE_CHAPTER = re.compile(r"^[" + CN_NUM + r"]+、")
RE_SECTION = re.compile(r"^（[" + CN_NUM + r"]+）")


def cn_quotes(text):
    out, open_q = [], True
    for ch in text:
        if ch == '"':
            out.append("\u201c" if open_q else "\u201d")
            open_q = not open_q
        else:
            out.append(ch)
    return "".join(out)


def set_font(run, name="宋体", size=12, bold=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_body(text, indent=True, size=12, bold=False, hanging=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    if hanging:
        pf.left_indent = Pt(24)
        pf.first_line_indent = Pt(-12)
    elif indent:
        pf.first_line_indent = Pt(size * 2)
    run = p.add_run(cn_quotes(text))
    set_font(run, size=size, bold=bold)
    return p


def add_heading(text, level):
    sizes = {1: 18, 2: 15, 3: 13}
    p = doc.add_paragraph(style=f"Heading {level}")
    pf = p.paragraph_format
    pf.space_before = Pt(18 if level <= 2 else 12)
    pf.space_after = Pt(10 if level <= 2 else 6)
    pf.line_spacing = 1.4
    if level == 1:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(cn_quotes(text))
    set_font(run, name="黑体", size=sizes.get(level, 12), bold=True,
             color=(0x1F, 0x37, 0x64) if level <= 2 else (0x2E, 0x2E, 0x2E))
    return p


with open(SRC, encoding="utf-8") as f:
    lines = f.read().splitlines()

i, n, front_matter = 0, len(lines), True

while i < n and not lines[i].strip():
    i += 1
if i < n:
    add_heading(lines[i].strip(), 1)
    i += 1
if i < n and lines[i].strip().startswith("（"):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(16)
    set_font(p.add_run(lines[i].strip()), name="黑体", size=13, color=(0x40, 0x40, 0x40))
    i += 1

body_count = 0

while i < n:
    raw = lines[i]
    stripped = raw.strip()

    if not stripped:
        i += 1
        continue

    if re.fullmatch(r"[─\-=]{4,}", stripped):
        i += 1
        continue

    leading = len(raw) - len(raw.lstrip())

    if leading == 0 and RE_CHAPTER.match(stripped):
        add_heading(stripped, 2)
        front_matter = False
        i += 1
        continue

    if leading == 0 and RE_SECTION.match(stripped):
        add_heading(stripped, 3)
        i += 1
        continue

    m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
    if m:
        add_heading(m.group(2).strip(), len(m.group(1)))
        i += 1
        continue

    if front_matter:
        add_body(stripped, indent=False, size=11)
        i += 1
        continue

    buf = [stripped]
    base = leading
    i += 1
    while i < n:
        nraw = lines[i]
        nxt = nraw.strip()
        if not nxt or re.fullmatch(r"[─\-=]{4,}", nxt):
            break
        nind = len(nraw) - len(nraw.lstrip())
        if nind == 0 and (RE_CHAPTER.match(nxt) or RE_SECTION.match(nxt)
                          or re.match(r"^#{1,4}\s+", nxt) or nxt.startswith(">")):
            break
        if nind <= base and re.match(r"^[-*]\s+", nxt):
            break
        if buf[-1].endswith(("：", ":")):
            break
        buf.append(nxt)
        i += 1

    text = ""
    for seg in buf:
        if text and re.search(r"[A-Za-z0-9)\]]$", text) and re.match(r"^[A-Za-z0-9(\[]", seg):
            text += " " + seg
        else:
            text += seg
    text = re.sub(r"\s{2,}", " ", text).strip()

    if 0 < leading <= 2 and RE_CHAPTER.match(text) and len(buf) == 1 and len(text) < 60:
        add_body(text, hanging=True)
        body_count += 1
        continue

    add_body(text)
    body_count += 1

doc.save(DST)
print(f"OK -> {DST} (正文段 {body_count})")
